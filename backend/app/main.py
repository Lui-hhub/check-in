import json
import io
import logging
import math
import time
import uuid
import asyncio
from collections import defaultdict, deque
from datetime import datetime
from contextlib import asynccontextmanager
from pathlib import Path
from typing import Annotated
from fastapi import Depends, FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse
from starlette.concurrency import run_in_threadpool
from PIL import Image, UnidentifiedImageError
from openpyxl import load_workbook
from openpyxl import Workbook
from sqlalchemy import desc, func, select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload
from .auth import create_token, require_any, require_role
from .config import get_settings
from .db import get_db
from .models import Checkin, Student, StudentSubject
from .schemas import CheckinResponse, LoginRequest, StudentFields, StudentImportResponse, StudentResponse, TokenResponse

settings = get_settings()
logging.basicConfig(
    level=getattr(logging, settings.log_level.upper(), logging.INFO),
    format="%(asctime)s %(levelname)s %(name)s %(message)s",
)
logger = logging.getLogger("checkin.api")
media_root = Path(settings.media_dir).resolve()

class InMemoryRateLimiter:
    """Per-process limiter; use Redis when the API runs on multiple hosts."""

    def __init__(self) -> None:
        self.requests: dict[tuple[str, str], deque[float]] = defaultdict(deque)
        self.lock = asyncio.Lock()

    async def allow(self, key: tuple[str, str], limit: int, now: float) -> bool:
        async with self.lock:
            bucket = self.requests[key]
            cutoff = now - 60
            while bucket and bucket[0] <= cutoff:
                bucket.popleft()
            if len(bucket) >= limit:
                return False
            bucket.append(now)
            return True

rate_limiter = InMemoryRateLimiter()

@asynccontextmanager
async def lifespan(_: FastAPI):
    media_root.mkdir(parents=True, exist_ok=True)
    yield

app = FastAPI(title="补习班签到 API", lifespan=lifespan)
allowed_origins = [origin.strip().rstrip("/") for origin in settings.frontend_origin.split(",") if origin.strip()]
app.add_middleware(CORSMiddleware, allow_origins=allowed_origins, allow_credentials=True, allow_methods=["*"], allow_headers=["*"])

@app.middleware("http")
async def rate_limit_requests(request: Request, call_next):
    if request.method == "OPTIONS" or not request.url.path.startswith("/api/"):
        return await call_next(request)
    path = request.url.path
    if path.startswith("/api/auth/"):
        limit = settings.rate_limit_auth_per_minute
        category = "auth"
    elif path == "/api/checkins":
        limit = settings.rate_limit_checkins_per_minute
        category = "checkins"
    else:
        limit = settings.rate_limit_api_per_minute
        category = "api"
    client = request.headers.get("x-real-ip") or (request.client.host if request.client else "unknown")
    if not await rate_limiter.allow((client, category), limit, time.monotonic()):
        logger.warning("rate limit exceeded category=%s client=%s", category, client)
        return JSONResponse({"detail": "请求过于频繁，请稍后再试"}, status_code=429, headers={"Retry-After": "60"})
    return await call_next(request)

@app.middleware("http")
async def request_logging(request: Request, call_next):
    started = time.perf_counter()
    client = request.client.host if request.client else "unknown"
    try:
        response = await call_next(request)
    except Exception:
        logger.exception("request failed method=%s path=%s client=%s", request.method, request.url.path, client)
        raise
    elapsed_ms = (time.perf_counter() - started) * 1000
    logger.info("request method=%s path=%s status=%s duration_ms=%.1f client=%s", request.method, request.url.path, response.status_code, elapsed_ms, client)
    return response

@app.get("/health")
async def health(): return {"status": "ok"}

@app.post("/api/auth/{kind}-login", response_model=TokenResponse)
async def login(kind: str, body: LoginRequest):
    passwords = {"viewer": settings.viewer_password, "uploader": settings.uploader_password, "admin": settings.admin_password}
    if kind not in passwords or body.password != passwords[kind]:
        logger.warning("login failed role=%s", kind)
        raise HTTPException(status_code=401, detail="密码错误")
    logger.info("login succeeded role=%s", kind)
    return TokenResponse(access_token=create_token(kind), role=kind)

def student_out(s: Student) -> StudentResponse:
    return StudentResponse(id=s.id, grade=s.grade, name=s.name, subjects=[item.subject for item in s.subjects], display_name=s.display_name, is_deleted=s.is_deleted)

async def unique_display_name(db: AsyncSession, grade: str, name: str, exclude_id: int | None = None) -> str:
    query = select(Student).where(Student.grade == grade, Student.name == name)
    if exclude_id is not None: query = query.where(Student.id != exclude_id)
    result = await db.execute(query)
    existing = list(result.scalars())
    if not existing: return name
    return f"{name}-{max([int(s.display_name.rsplit('-', 1)[1]) for s in existing if s.display_name.startswith(name + '-') and s.display_name.rsplit('-', 1)[1].isdigit()] or [1]) + 1}"

async def next_student_id(db: AsyncSession) -> int:
    year = datetime.now().year
    start, end = year * 10000, year * 10000 + 9999
    value = await db.scalar(select(func.max(Student.id)).where(Student.id >= start, Student.id <= end))
    if value is not None and value >= end:
        raise HTTPException(409, "本年度学生编号已用完")
    return (value or start) + 1

def save_image(data: bytes, target: Path) -> None:
    """Decode and normalize an uploaded image outside the async event loop."""
    try:
        with Image.open(io.BytesIO(data)) as image:
            image.load()
            if image.width > 640 or image.height > 640:
                raise HTTPException(400, "上传照片尺寸不能超过 640 像素")
            image.convert("RGB").save(target, format="JPEG", quality=80, optimize=True)
    except (UnidentifiedImageError, OSError) as exc:
        raise HTTPException(400, "请上传有效的 JPEG 照片") from exc

async def save_upload(upload: UploadFile, folder: str) -> str:
    suffix = ".jpg"
    relative = f"{folder}/{uuid.uuid4().hex}{suffix}"
    target = media_root / relative
    target.parent.mkdir(parents=True, exist_ok=True)
    data = await upload.read()
    if len(data) > 8 * 1024 * 1024: raise HTTPException(413, "图片过大")
    await run_in_threadpool(save_image, data, target)
    return relative

def parse_embedding(raw: str) -> list[float]:
    try: value = json.loads(raw)
    except json.JSONDecodeError as exc: raise HTTPException(400, "特征格式错误") from exc
    if not isinstance(value, list) or len(value) != 128 or not all(isinstance(x, (int, float)) for x in value): raise HTTPException(400, "特征必须是 128 维数组")
    return [float(x) for x in value]

def distance(a: list[float], b: list[float]) -> float:
    return math.sqrt(sum((x - y) ** 2 for x, y in zip(a, b)))

@app.get("/api/students", response_model=list[StudentResponse])
async def list_students(include_deleted: bool = False, _: dict = Depends(require_any("uploader", "admin")), db: AsyncSession = Depends(get_db)):
    query = select(Student).options(selectinload(Student.subjects)).order_by(Student.grade, Student.display_name)
    if not include_deleted: query = query.where(Student.is_deleted.is_(False))
    return [student_out(s) for s in (await db.execute(query)).scalars()]

async def create_student(db: AsyncSession, grade: str, name: str, subject: str | list[str], embedding: list[float], image_path: str | None):
    subjects = [subject] if isinstance(subject, str) else subject
    student = Student(id=await next_student_id(db), grade=grade, name=name, display_name=await unique_display_name(db, grade, name), face_embedding=embedding, face_embedding_vector=embedding, face_image_path=image_path)
    student.subjects = [StudentSubject(subject=item.strip()) for item in subjects if item.strip()]
    db.add(student); await db.commit(); await db.refresh(student); return student

@app.post("/api/students", response_model=StudentResponse)
async def add_student(grade: Annotated[str, Form()], name: Annotated[str, Form()], subjects: Annotated[str, Form()], embedding: Annotated[str, Form()], face_image: Annotated[UploadFile, File()], _: dict = Depends(require_role("admin")), db: AsyncSession = Depends(get_db)):
    path = await save_upload(face_image, "students")
    subject_list = [item.strip() for item in subjects.replace("，", ",").split(",") if item.strip()]
    if not subject_list: raise HTTPException(400, "至少填写一个科目")
    student = await create_student(db, grade, name, subject_list, parse_embedding(embedding), path)
    logger.info("student created student_id=%s grade=%s subjects=%s", student.id, grade, ",".join(subject_list))
    return student_out(student)

@app.post("/api/students/import", response_model=StudentImportResponse)
async def import_students(file: Annotated[UploadFile, File()], _: dict = Depends(require_role("admin")), db: AsyncSession = Depends(get_db)):
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(400, "请上传 .xlsx 格式的 Excel 文件")
    data = await file.read()
    if len(data) > 5 * 1024 * 1024:
        raise HTTPException(413, "Excel 文件不能超过 5MB")
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, None)
        if not headers:
            raise HTTPException(400, "Excel 文件没有表头")
        aliases = {"年级": "grade", "grade": "grade", "姓名": "name", "name": "name", "补课科目": "subject", "科目": "subject", "subject": "subject"}
        columns = {aliases.get(str(value).strip().lower()): index for index, value in enumerate(headers) if value is not None and aliases.get(str(value).strip().lower())}
        missing = [label for label, key in (("年级", "grade"), ("姓名", "name"), ("补课科目", "subject")) if key not in columns]
        if missing:
            raise HTTPException(400, f"Excel 缺少列：{', '.join(missing)}")
        created = 0
        errors: list[str] = []
        for row_number, row in enumerate(rows, start=2):
            values = {key: str(row[index]).strip() if index < len(row) and row[index] is not None else "" for key, index in columns.items()}
            if not any(values.values()):
                continue
            missing_values = [label for label, key in (("年级", "grade"), ("姓名", "name"), ("补课科目", "subject")) if not values[key]]
            if missing_values:
                errors.append(f"第 {row_number} 行缺少：{', '.join(missing_values)}")
                continue
            student = Student(id=await next_student_id(db), grade=values["grade"], name=values["name"], display_name=await unique_display_name(db, values["grade"], values["name"]), face_embedding=[], face_embedding_vector=None, face_image_path=None)
            student.subjects = [StudentSubject(subject=item.strip()) for item in values["subject"].replace("，", ",").split(",") if item.strip()]
            db.add(student)
            await db.flush()
            created += 1
        await db.commit()
        workbook.close()
        logger.info("students imported created=%s errors=%s", created, len(errors))
        return StudentImportResponse(created=created, errors=errors)
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(400, "无法读取 Excel 文件，请确认文件未损坏且为 .xlsx 格式") from exc

@app.get("/api/students/export")
async def export_students(include_deleted: bool = False, ids: str | None = None, _: dict = Depends(require_role("admin")), db: AsyncSession = Depends(get_db)):
    query = select(Student).options(selectinload(Student.subjects)).order_by(Student.grade, Student.display_name)
    if include_deleted:
        query = query.where(Student.is_deleted.is_(True))
    else:
        query = query.where(Student.is_deleted.is_(False))
    if ids:
        try:
            selected = [int(value) for value in ids.split(",") if value]
        except ValueError as exc:
            raise HTTPException(400, "学生编号格式错误") from exc
        query = query.where(Student.id.in_(selected))
    students = list((await db.execute(query)).scalars())
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "学生名单"
    sheet.append(["学生编号", "年级", "姓名", "补课科目", "状态"])
    for student in students:
        sheet.append([student.id, student.grade, student.display_name, "、".join(item.subject for item in student.subjects), "毕业生" if student.is_deleted else "在读"])
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": "attachment; filename=students.xlsx"})

@app.patch("/api/students/{student_id}", response_model=StudentResponse)
async def edit_student(student_id: int, fields: StudentFields, _: dict = Depends(require_role("admin")), db: AsyncSession = Depends(get_db)):
    student = await db.scalar(select(Student).options(selectinload(Student.subjects)).where(Student.id == student_id))
    if not student: raise HTTPException(404, "学生不存在")
    changed_identity = (student.grade, student.name) != (fields.grade, fields.name)
    student.grade, student.name = fields.grade, fields.name
    if changed_identity:
        student.display_name = await unique_display_name(db, fields.grade, fields.name, exclude_id=student.id)
    student.subjects.clear()
    student.subjects.extend(StudentSubject(subject=item.strip()) for item in fields.subjects if item.strip())
    await db.commit(); await db.refresh(student)
    logger.info("student updated student_id=%s grade=%s subjects=%s", student.id, student.grade, ",".join(fields.subjects))
    return student_out(student)

@app.post("/api/students/{student_id}/face", response_model=StudentResponse)
async def replace_face(student_id: int, embedding: Annotated[str, Form()], face_image: Annotated[UploadFile, File()], _: dict = Depends(require_role("admin")), db: AsyncSession = Depends(get_db)):
    student = await db.scalar(select(Student).options(selectinload(Student.subjects)).where(Student.id == student_id))
    if not student: raise HTTPException(404, "学生不存在")
    student.face_embedding = parse_embedding(embedding); student.face_embedding_vector = student.face_embedding; student.face_image_path = await save_upload(face_image, "students")
    await db.commit(); await db.refresh(student)
    logger.info("student face updated student_id=%s", student.id)
    return student_out(student)

@app.delete("/api/students/{student_id}")
async def delete_student(student_id: int, _: dict = Depends(require_role("admin")), db: AsyncSession = Depends(get_db)):
    student = await db.get(Student, student_id)
    if not student: raise HTTPException(404, "学生不存在")
    student.is_deleted = True; await db.commit()
    logger.info("student soft deleted student_id=%s", student.id)
    return {"ok": True}

@app.post("/api/checkins", response_model=CheckinResponse)
async def add_checkin(embedding: Annotated[str, Form()], photo: Annotated[UploadFile, File()], student_id: Annotated[int | None, Form()] = None, subject: Annotated[str | None, Form()] = None, _: dict = Depends(require_role("uploader")), db: AsyncSession = Depends(get_db)):
    vector = parse_embedding(embedding)
    student = await db.scalar(select(Student).options(selectinload(Student.subjects)).where(Student.id == student_id)) if student_id else None
    if student and student.is_deleted: raise HTTPException(422, "学生已删除")
    matched_distance = None
    if not student:
        candidates = list((await db.execute(
            select(Student.id, Student.face_embedding_vector.l2_distance(vector).label("distance"))
            .where(Student.is_deleted.is_(False), Student.face_embedding_vector.is_not(None))
            .order_by(Student.face_embedding_vector.l2_distance(vector))
            .limit(2)
        )).all())
        if not candidates: raise HTTPException(422, "暂无学生资料")
        if len(candidates) > 1 and float(candidates[1].distance) - float(candidates[0].distance) <= 0.02:
            candidate_ids = [row.id for row in candidates]
            candidate_students = list((await db.execute(
                select(Student).options(selectinload(Student.subjects)).where(Student.id.in_(candidate_ids))
            )).scalars())
            by_id = {item.id: item for item in candidate_students}
            raise HTTPException(409, detail={
                "code": "STUDENT_AMBIGUOUS",
                "message": "照片对应多个相似的学生资料，请手动选择学生",
                "candidates": [
                    {"student_id": row.id, "display_name": by_id[row.id].display_name, "subjects": [item.subject for item in by_id[row.id].subjects]}
                    for row in candidates if row.id in by_id
                ],
            })
        matched_id, matched_distance = candidates[0].id, float(candidates[0].distance)
        student = await db.scalar(
            select(Student).options(selectinload(Student.subjects)).where(Student.id == matched_id)
        )
        if student is None: raise HTTPException(422, "学生资料已变更，请重试")
        logger.info("face match candidate student_id=%s distance=%.4f threshold=%.4f", student.id, matched_distance, settings.face_match_threshold)
        if matched_distance > settings.face_match_threshold:
            logger.warning("face match failed distance=%.4f threshold=%.4f", matched_distance, settings.face_match_threshold)
            raise HTTPException(422, "未识别到学生")
    else:
        logger.info("manual student selected student_id=%s", student.id)
    if student.is_deleted: raise HTTPException(422, "学生已删除")
    available_subjects = [item.subject for item in student.subjects]
    if subject not in available_subjects:
        if subject is None and len(available_subjects) == 1: subject = available_subjects[0]
        else:
            raise HTTPException(409, detail={
                "code": "SUBJECT_REQUIRED",
                "message": "请选择该学生的补课科目",
                "student_id": student.id,
                "display_name": student.display_name,
                "subjects": available_subjects,
                "distance": matched_distance,
            })
    path = await save_upload(photo, "checkins")
    checkin = Checkin(student_id=student.id, photo_path=path, grade_snapshot=student.grade, name_snapshot=student.display_name, subject_snapshot=subject)
    db.add(checkin); await db.flush()
    old = list((await db.execute(select(Checkin).where(Checkin.student_id == student.id).order_by(desc(Checkin.created_at), desc(Checkin.id)))).scalars())
    for record in old[10:]:
        try: (media_root / record.photo_path).unlink(missing_ok=True)
        except OSError: pass
        await db.delete(record)
    await db.commit(); await db.refresh(checkin)
    logger.info("checkin created checkin_id=%s student_id=%s distance=%s", checkin.id, student.id, f"{matched_distance:.4f}" if matched_distance is not None else "manual")
    return CheckinResponse(id=checkin.id, student_id=student.id, photo_url=f"/media/{path}", grade=checkin.grade_snapshot, name=checkin.name_snapshot, subject=checkin.subject_snapshot, created_at=checkin.created_at, distance=matched_distance)

@app.get("/api/checkins", response_model=list[CheckinResponse])
async def list_checkins(_: dict = Depends(require_any("viewer", "admin")), db: AsyncSession = Depends(get_db)):
    records = list((await db.execute(select(Checkin).join(Checkin.student).where(Student.is_deleted.is_(False)).order_by(desc(Checkin.created_at), desc(Checkin.id)))).scalars())
    return [CheckinResponse(id=r.id, student_id=r.student_id, photo_url=f"/media/{r.photo_path}", grade=r.grade_snapshot, name=r.name_snapshot, subject=r.subject_snapshot, created_at=r.created_at) for r in records]

@app.get("/media/{path:path}")
async def media(path: str, _: dict = Depends(require_any("viewer", "uploader", "admin"))):
    target = (media_root / path).resolve()
    if media_root not in target.parents or not target.is_file(): raise HTTPException(404, "文件不存在")
    return FileResponse(target)
